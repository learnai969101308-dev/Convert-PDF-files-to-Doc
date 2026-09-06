import os
import re
from flask import Flask, render_template, request, send_file
from pdf2image import convert_from_path
import pytesseract
from docx import Document
from pptx import Presentation
import pandas as pd
import cv2
import numpy as np
from PIL import Image

app = Flask(__name__)
UPLOAD_FOLDER, OUTPUT_FOLDER = 'uploads', 'outputs'
os.makedirs(UPLOAD_FOLDER, exist_ok=True)
os.makedirs(OUTPUT_FOLDER, exist_ok=True)

def remove_background_opencv(image):
    open_cv_image = np.array(image)
    gray = cv2.cvtColor(open_cv_image, cv2.COLOR_RGB2GRAY)
    filtered = cv2.bilateralFilter(gray, 9, 75, 75)
    thresh = cv2.adaptiveThreshold(filtered, 255, cv2.ADAPTIVE_THRESH_GAUSSIAN_C, cv2.THRESH_BINARY, 11, 12)
    return Image.fromarray(thresh)

def clean_khmer_text(text):
    return '\n'.join([line.strip() for line in re.sub(r' +', ' ', text if text else "").split('\n') if line.strip()])

@app.route('/')
def index():
    return render_template('index.html')

@app.route('/convert', methods=['POST'])
def convert():
    if 'file' not in request.files or request.files['file'].filename == '':
        return "គ្មានឯកសារ", 400
    
    file = request.files['file']
    conversion_type = request.form.get('type')
    file_path = os.path.join(UPLOAD_FOLDER, file.filename)
    file.save(file_path)
    base_name = file.filename.rsplit('.', 1)[0]
    
    try:
        images = convert_from_path(file_path, 150)

        if conversion_type == 'word':
            docx_path = os.path.join(OUTPUT_FOLDER, base_name + '.docx')
            doc = Document()
            for idx, img in enumerate(images):
                text = clean_khmer_text(pytesseract.image_to_string(remove_background_opencv(img), lang='khm'))
                if text: doc.add_paragraph(text)
            doc.save(docx_path)
            return send_file(docx_path, as_attachment=True)
            
        elif conversion_type == 'excel':
            excel_path = os.path.join(OUTPUT_FOLDER, base_name + '.xlsx')
            all_rows = []
            
            for page_num, img in enumerate(images):
                df_data = pytesseract.image_to_data(remove_background_opencv(img), lang='khm', output_type=pytesseract.Output.DATAFRAME)
                df_data = df_data[df_data.text.notnull() & (df_data.text.str.strip() != '')]
                
                if not df_data.empty:
                    median_height = df_data['height'].median()
                    y_tolerance = max(10, median_height * 0.6) 
                    x_tolerance = max(20, median_height * 1.5) 
                    
                    df_data = df_data.sort_values('top')
                    rows_grouped, current_row, last_top = [], [], None
                    
                    for _, row in df_data.iterrows():
                        if last_top is None:
                            current_row.append(row)
                            last_top = row['top']
                        else:
                            avg_top = sum([r['top'] for r in current_row]) / len(current_row)
                            if abs(row['top'] - avg_top) <= y_tolerance:
                                current_row.append(row)
                            else:
                                rows_grouped.append(pd.DataFrame(current_row))
                                current_row = [row]
                                last_top = row['top']
                    if current_row: rows_grouped.append(pd.DataFrame(current_row))
                    
                    for r_df in rows_grouped:
                        r_df = r_df.sort_values('left')
                        cols, current_col_text, last_right = [], [], -1
                        
                        for _, row in r_df.iterrows():
                            if last_right == -1:
                                current_col_text.append(str(row['text']))
                                last_right = row['left'] + row['width']
                            else:
                                if row['left'] - last_right > x_tolerance: 
                                    cols.append(" ".join(current_col_text))
                                    current_col_text = [str(row['text'])]
                                else:
                                    current_col_text.append(str(row['text']))
                                last_right = row['left'] + row['width']
                                
                        if current_col_text: cols.append(" ".join(current_col_text))
                        if cols: all_rows.append({'ទំព័រទី': page_num + 1, **{f'ជួរឈរ_{i+1}': c for i, c in enumerate(cols)}})
            
            pd.DataFrame(all_rows).to_excel(excel_path, index=False)
            return send_file(excel_path, as_attachment=True)
            
        elif conversion_type == 'pptx':
            pptx_path = os.path.join(OUTPUT_FOLDER, base_name + '.pptx')
            prs = Presentation()
            for idx, img in enumerate(images):
                text = clean_khmer_text(pytesseract.image_to_string(remove_background_opencv(img), lang='khm'))
                slide = prs.slides.add_slide(prs.slide_layouts[1])
                slide.shapes.title.text = f"ទំព័រទី {idx + 1}"
                slide.placeholders[1].text = text if text else "គ្មានទិន្នន័យ"
            prs.save(pptx_path)
            return send_file(pptx_path, as_attachment=True)
            
    except Exception as e:
        return f"កំហុស៖ {e}", 500

@app.route('/process_image', methods=['POST'])
def process_image():
    if 'image_file' not in request.files:
        return "គ្មានរូបភាព", 400
    
    file = request.files['image_file']
    action = request.form.get('action', 'clean') 
    
    x = int(float(request.form.get('x', 0)))
    y = int(float(request.form.get('y', 0)))
    w = int(float(request.form.get('w', 0)))
    h = int(float(request.form.get('h', 0)))
    
    img_path = os.path.join(UPLOAD_FOLDER, file.filename)
    file.save(img_path)
    
    image = cv2.imread(img_path)
    if w > 0 and h > 0:
        image = image[y:y+h, x:x+w]
        
    processed_img = remove_background_opencv(Image.fromarray(cv2.cvtColor(image, cv2.COLOR_BGR2RGB)))
    
    if action == 'word':
        docx_path = os.path.join(OUTPUT_FOLDER, 'Image_Result.docx')
        doc = Document()
        doc.add_heading('អត្ថបទស្រង់ចេញពីរូបភាព', level=1)
        
        text = clean_khmer_text(pytesseract.image_to_string(processed_img, lang='khm'))
        if text:
            doc.add_paragraph(text)
        else:
            doc.add_paragraph("មិនមានអត្ថបទត្រូវបានរកឃើញនៅក្នុងរូបភាពនេះទេ!")
            
        doc.save(docx_path)
        return send_file(docx_path, as_attachment=True)
        
    else:
        output_img_path = os.path.join(OUTPUT_FOLDER, 'HD_Cleaned_Image.png')
        processed_img.save(output_img_path)
        return send_file(output_img_path, as_attachment=True)

# ចំណុចផ្លាស់ប្តូរសម្រាប់ GitHub / Server Deployment
if __name__ == '__main__':
    # ប្រើ host='0.0.0.0' ដើម្បីឱ្យ Server (ដូចជា Heroku/Render) អាចភ្ជាប់បាន
    app.run(host='0.0.0.0', port=int(os.environ.get('PORT', 5000)), debug=False)